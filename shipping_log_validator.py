"""DT production writing-log analysis; independent of shipment SN/MAC ranges."""
from __future__ import annotations
import datetime, logging, re, secrets
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from shipping_validator import DEFAULTS, pass_rules

_FIELDS = {
 'ssid_mes':r'\[MES\]\s*字段查询响应：[^\n]*?FieldValue=([^\r\n,]+)',
 'ssid_set':r'命令：nvram set WiFiSSIDSuffix\s+(\S+)',
 'ssid_get':r'命令：nvram get WiFiSSIDSuffix\s*\r?\n\s*stdout：[^\r\n]*\r?\n\s*(\S+)',
 'key_mes':r'MES 字段“cst10”读取成功，值：([^。\r\n]+)',
 'key_set':r'命令：nvram set WiFiPwd\s+(\S+)',
 'key_get':r'命令：nvram get WiFiPwd\s*\r?\n\s*stdout：[^\r\n]*\r?\n\s*(\S+)',
 'serial':r'命令：nvram set SerialNum\s+(\S+)',
 'mac':r'输入值“mac”已确认：([0-9A-Fa-f]{12})',
 'final':r'\[系统\] 最终结果：([^\r\n]+)',
}

def decode_file(path):
 raw=Path(path).read_bytes()
 for encoding in ('utf-8-sig','utf-16','gb18030'):
  try:return raw.decode(encoding),encoding
  except UnicodeDecodeError:continue
 raise ValueError('Cannot decode log as UTF-8, UTF-16 or GB18030')

def check_log(text,cfg=None):
 cfg={**DEFAULTS,**(cfg or {})}; matches={}
 for name,regex in _FIELDS.items():
  found=list(re.finditer(regex,text,re.MULTILINE))
  matches[name]=[(m.group(1).strip(),text.count('\n',0,m.start())+1) for m in found]
 issues=[]
 def problem(field,code,detail,line=''):
  issues.append({'field':field,'code':code,'detail':detail,'line':line})
 def get(name,mandatory=True):
  entries=matches[name]
  if not entries:
   if mandatory:problem(name,'LOG_MISSING_FIELD','No evidence for '+name)
   return None
  values={v for v,_ in entries}
  if len(values)>1:problem(name,'LOG_CONFLICTING_VALUES','Multiple different values found',entries[0][1])
  return entries[-1][0]
 suffix_mes=get('ssid_mes');suffix_set=get('ssid_set');suffix_get=get('ssid_get')
 key_mes=get('key_mes');key_set=get('key_set');key_get=get('key_get')
 sn=get('serial',False);mac=get('mac',False);final=get('final',False)
 if suffix_mes is not None:
  prefix=cfg['ssid_24_prefix']
  if not suffix_mes.startswith(prefix):problem('SSID','SSID_MES_PREFIX','MES SSID does not start with '+prefix,matches['ssid_mes'][-1][1]);expected_suffix=None
  else:expected_suffix=suffix_mes[len(prefix):]
  if expected_suffix is not None and (len(expected_suffix)!=int(cfg['ssid_suffix_length']) or not re.fullmatch(cfg['ssid_suffix_regex'],expected_suffix,re.ASCII)):
   problem('SSID','SSID_SUFFIX_FORMAT','MES suffix must match configured length and pattern',matches['ssid_mes'][-1][1])
  if suffix_set is not None and expected_suffix is not None and suffix_set!=expected_suffix:problem('SSID','SSID_WRITE_MISMATCH','Written suffix differs from MES suffix',matches['ssid_set'][-1][1])
  if suffix_get is not None and expected_suffix is not None and suffix_get!=expected_suffix:problem('SSID','SSID_READBACK_MISMATCH','Readback suffix differs from MES suffix',matches['ssid_get'][-1][1])
 if suffix_set is not None and suffix_get is not None and suffix_set!=suffix_get:problem('SSID','SSID_SET_GET_MISMATCH','Written and readback suffix differ',matches['ssid_get'][-1][1])
 if key_mes is not None:
  for err in pass_rules(key_mes,cfg):problem('WiFi WPA Key',err,'MES WiFi key violates '+err,matches['key_mes'][-1][1])
  if key_set is not None and key_set!=key_mes:problem('WiFi WPA Key','KEY_WRITE_MISMATCH','Written key differs from MES source',matches['key_set'][-1][1])
  if key_get is not None and key_get!=key_mes:problem('WiFi WPA Key','KEY_READBACK_MISMATCH','Readback differs from MES source',matches['key_get'][-1][1])
 if key_set is not None and key_get is not None and key_set!=key_get:problem('WiFi WPA Key','KEY_SET_GET_MISMATCH','Written and readback keys differ',matches['key_get'][-1][1])
 if final is not None and '成功' not in final:problem('Production test','PRODUCTION_NOT_PASS','Final station status is not success',matches['final'][-1][1])
 return {'values':{'SN':sn,'MAC':mac,'MES 2.4G SSID':suffix_mes,'expected 5G SSID':(cfg['ssid_5_prefix']+suffix_mes[len(cfg['ssid_24_prefix']):] if suffix_mes and suffix_mes.startswith(cfg['ssid_24_prefix']) else ''),'written SSID suffix':suffix_set,'readback SSID suffix':suffix_get,'MES WPA Key':key_mes,'written WPA Key':key_set,'readback WPA Key':key_get,'station final':final},'issues':issues,'status':'FAIL' if issues else 'PASS','evidence':matches,'note':'5G complete SSID inferred from configured prefix; log only writes/reads common SSID suffix. Actual 5G broadcast SSID not verified.'}

def analyze_log(path,output_dir,cfg=None):
 cfg={**DEFAULTS,**(cfg or {})}; out=Path(output_dir);out.mkdir(parents=True,exist_ok=True)
 run=datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+'_'+secrets.token_hex(3)
 debug=out/f'DEBUG_LOG_{run}.log';execution=out/f'EXECUTION_LOG_{run}.log';report=out/f'Log_Validation_{run}.xlsx'
 try:
  text,encoding=decode_file(path);result=check_log(text,cfg)
  with open(execution,'w',encoding='utf-8') as f:
   f.write(f'Version 1.2.0 | Source {Path(path).name} | Encoding {encoding}\nResult {result["status"]} | Issues {len(result["issues"])}\n')
   for issue in result['issues']:f.write(f'{issue["code"]} line={issue["line"]}\n')
  with open(debug,'w',encoding='utf-8') as f:
   f.write('Extracted evidence (secrets masked):\n')
   for name,entries in result['evidence'].items():
    f.write(f'{name}: '+', '.join(f'line {line} (length {len(value)})' if 'key' in name else f'line {line} {value}' for value,line in entries)+'\n')
  wb=Workbook();ws=wb.active;ws.title='Summary'
  for row in [('DT log validation','V1.2.0'),('File',Path(path).name),('Encoding',encoding),('Result',result['status']),('Number of issues',len(result['issues'])),('SSID policy','Repeated SSIDs allowed'),('Note',result['note']),('Policy boundary','This log is checked independently of shipment SN/MAC ranges')]:ws.append(row)
  s=wb.create_sheet('Fields');s.append(['Field','Extracted value','Evidence line'])
  field_to_log={'SN':'serial','MAC':'mac','MES 2.4G SSID':'ssid_mes','written SSID suffix':'ssid_set','readback SSID suffix':'ssid_get','MES WPA Key':'key_mes','written WPA Key':'key_set','readback WPA Key':'key_get','station final':'final'}
  for k,v in result['values'].items():s.append([k,v,', '.join(str(line) for _,line in result['evidence'].get(field_to_log.get(k,''),[]))])
  fail=wb.create_sheet('Failures');fail.append(['Field','Code','Detail','Log line'])
  for i in result['issues']:fail.append([i['field'],i['code'],i['detail'],i['line']])
  rules=wb.create_sheet('Rules');rules.append(['Rule','Value'])
  for k in ('ssid_24_prefix','ssid_5_prefix','ssid_suffix_length','ssid_suffix_regex','password_length','allowed_special','banned'):rules.append([k,str(cfg[k])])
  for sheet in wb:
   sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
   for c in sheet[1]:c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(color='FFFFFF',bold=True)
   for letter,width in [('A',30),('B',68),('C',33),('D',24)]:sheet.column_dimensions[letter].width=width
  try:wb.save(report)
  finally:wb.close()
  return {'report':str(report),'execution_log':str(execution),'debug_log':str(debug),'pass':int(result['status']=='PASS'),'fail':int(result['status']=='FAIL'),'errors':{i['code']:1 for i in result['issues']},'total':1,'details':result}
 except Exception as e:
  with open(debug,'a',encoding='utf-8') as f:f.write(f'FATAL: {type(e).__name__}: {e}\n')
  raise
