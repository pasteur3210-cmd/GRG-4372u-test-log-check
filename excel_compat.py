"""Small compatibility facade using the public PyPI openpyxl dependency.
Replaces internal-only artifact_tool workbook interface for Windows distribution.
"""
from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils.cell import range_boundaries, get_column_letter

class Format:
    def __init__(self, sheet, cells): self.sheet=sheet;self.cells=cells
    @property
    def column_width(self): return None
    @column_width.setter
    def column_width(self, width):
        for col in sorted({c.column for row in self.cells for c in row}):
            self.sheet.column_dimensions[get_column_letter(col)].width = width
    @property
    def row_height(self): return None
    @row_height.setter
    def row_height(self, height):
        for row in sorted({c.row for cells in self.cells for c in cells}):
            self.sheet.row_dimensions[row].height=height

class Range:
    def __init__(self,sheet,cells):self.sheet=sheet;self.cells=cells
    @property
    def values(self):return [[c.value for c in row] for row in self.cells]
    @values.setter
    def values(self,values):
        for row, vals in zip(self.cells, values):
            for cell, value in zip(row,vals):cell.value=value
    @property
    def format(self):return Format(self.sheet,self.cells)
    @format.setter
    def format(self, specs):
        fill=specs.get('fill');font=specs.get('font',{})
        for row in self.cells:
            for c in row:
                if fill:c.fill=PatternFill('solid',fgColor=fill.lstrip('#'))
                if font:c.font=Font(bold=font.get('bold',False),color=font.get('color','#000000').lstrip('#'),size=font.get('size',11))
        if 'row_height' in specs:self.format.row_height=specs['row_height']
        if 'column_width' in specs:self.format.column_width=specs['column_width']
class Freeze:
    def __init__(self,sheet):self.sheet=sheet
    def freeze_rows(self,n):self.sheet.freeze_panes=f'A{n+1}'
class Sheet:
    def __init__(self,sheet):self.sheet=sheet;self.freeze_panes=Freeze(sheet)
    def get_range(self,reference):
        if ':' in reference and reference.split(':')[0].isalpha():
            start,end=reference.split(':');lo=self.sheet[start+'1'].column;hi=self.sheet[end+'1'].column
            return Range(self.sheet,[[self.sheet.cell(1,c) for c in range(lo,hi+1)]])
        minc,minr,maxc,maxr=range_boundaries(reference)
        return Range(self.sheet,[[self.sheet.cell(r,c) for c in range(minc,maxc+1)] for r in range(minr,maxr+1)])
    def get_range_by_indexes(self,row,col,height,width):
        return Range(self.sheet,[[self.sheet.cell(r,c) for c in range(col+1,col+width+1)] for r in range(row+1,row+height+1)])
class Sheets:
    def __init__(self,wb):self.wb=wb
    def add(self,name):return Sheet(self.wb.create_sheet(name))
class Workbook:
    @staticmethod
    def create():
        wb=OpenpyxlWorkbook();wb.remove(wb.active);return type('Book',(),{'worksheets':Sheets(wb),'_wb':wb})()
class SpreadsheetFile:
    @staticmethod
    def export_xlsx(wb):
        class Export:
            def save(self,path):wb._wb.save(path)
        return Export()
