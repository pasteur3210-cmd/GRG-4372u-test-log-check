import unittest, tempfile, pathlib, collections
from report import make_report
from inspector import CONFIG
from openpyxl import load_workbook

class TestV1511(unittest.TestCase):
 def test_classification_and_counts(self):
  states=['PASS','不符合','規格差異','FAIL','待確認','缺少證據']
  rows=[]
  for idx,status in enumerate(states):
   rows.append({'file':'test.log','category':'WiFi','item':'EVM','result':status,'actual':idx,'line':idx+1,'note':'test'})
  rows.append({'file':'test.log','category':'ParserStatus','item':'parser','result':'待確認','actual':'','line':0,'note':'parser'})
  with tempfile.TemporaryDirectory() as d:
   path=pathlib.Path(d)/'result.xlsx'
   make_report(rows,[],collections.Counter(x['result'] for x in rows if x['category']!='ParserStatus'),path,CONFIG)
   wb=load_workbook(path,read_only=True)
   try:
    self.assertEqual(wb['查核總覽']['B10'].value,1)
    self.assertTrue(set(states).issubset(wb.sheetnames))
    self.assertEqual(wb['逐筆比對'].max_row,7)
    for state in states:
     self.assertEqual(wb[state].max_row,2)
     self.assertEqual(wb[state]['A2'].value,'test.log')
    self.assertEqual(wb['程式解析紀錄'].max_row,2)
   finally:wb.close()
 def test_empty_sheets(self):
  with tempfile.TemporaryDirectory() as d:
   path=pathlib.Path(d)/'empty.xlsx';make_report([],[],{},path,CONFIG)
   wb=load_workbook(path,read_only=True)
   try:
    for state in ['PASS','不符合','規格差異','FAIL','待確認']:
     self.assertEqual(wb[state].max_row,1)
   finally:wb.close()
 def test_gui_does_not_autoload(self):
  src=pathlib.Path(__file__).with_name('enhanced.py').read_text(encoding='utf-8')
  self.assertNotIn('root.after(250,do_audit)',src)
  self.assertIn('return gui(None,args.output_dir,args.unit_ids)',src)
